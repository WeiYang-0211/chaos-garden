from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont

from app.extract import extract_bytes


def test_docx_paragraphs_and_table(tmp_path: Path):
    document = Document()
    document.add_paragraph("见字如面。")
    document.add_paragraph("第二段有 English words。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "手迹"
    buffer = BytesIO()
    document.save(buffer)
    result = extract_bytes("note.docx", buffer.getvalue())
    assert "见字如面" in result.text
    assert "English words" in result.text
    assert "姓名" in result.text


@pytest.mark.slow
def test_image_ocr(tmp_path: Path):
    pytest.importorskip("rapidocr_onnxruntime")
    font_path = Path(__file__).resolve().parents[1] / "assets" / "fonts" / "Xiaolai-Regular.ttf"
    image = Image.new("RGB", (1000, 240), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(str(font_path), 58)
    draw.text((40, 70), "见字如面，手迹测试。", font=font, fill=(20, 20, 20))
    path = tmp_path / "ocr.png"
    image.save(path)
    result = extract_bytes("ocr.png", path.read_bytes())
    assert "见字如面" in result.text
    assert result.source_type == "image"
