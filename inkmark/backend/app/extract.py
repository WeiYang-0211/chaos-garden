from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import re
import threading
from typing import Iterable

import numpy as np
from PIL import Image, ImageOps, UnidentifiedImageError

from .config import MAX_CHARS


class ExtractionError(ValueError):
    pass


@dataclass
class ExtractedDocument:
    text: str
    source_type: str
    title: str
    warnings: list[str] = field(default_factory=list)
    pages: int = 0

    def normalized(self) -> str:
        text = normalize_text(self.text)
        if len(text) > MAX_CHARS:
            self.warnings.append(f"文本超过 {MAX_CHARS:,} 字，已截取前 {MAX_CHARS:,} 字。")
            text = text[:MAX_CHARS]
        return text


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ").replace("\u3000", "　")
    text = "".join(char for char in text if char == "\n" or char == "\t" or ord(char) >= 32)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def _clean_title(name: str) -> str:
    stem = Path(name or "未命名手稿").stem.strip() or "未命名手稿"
    stem = re.sub(r"[\\/:*?\"<>|]+", " ", stem)
    return stem[:80]


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("无法识别文本编码，请另存为 UTF-8 文本后重试。")


def _iter_docx_blocks(document) -> Iterable[str]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if text:
                yield text
            else:
                yield ""
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows: list[str] = []
            for row in table.rows:
                cells: list[str] = []
                previous = None
                for cell in row.cells:
                    if cell._tc is previous:
                        continue
                    previous = cell._tc
                    cell_text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                    cells.append(cell_text)
                line = "　".join(cell for cell in cells if cell)
                if line.strip():
                    rows.append(line)
            if rows:
                yield "\n".join(rows)
                yield ""


def extract_docx(data: bytes, filename: str) -> ExtractedDocument:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError
    import zipfile

    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            uncompressed = sum(item.file_size for item in archive.infolist())
            if uncompressed > 250 * 1024 * 1024:
                raise ExtractionError("Word 文档解压后过大，已拒绝处理。")
    except zipfile.BadZipFile as exc:
        raise ExtractionError("Word 文件不是有效的 .docx 压缩包。") from exc

    try:
        document = Document(BytesIO(data))
    except (PackageNotFoundError, ValueError, KeyError) as exc:
        raise ExtractionError("Word 文件无法解析，可能已损坏或不是 .docx 格式。") from exc

    blocks = [block for block in _iter_docx_blocks(document) if block is not None]
    text = "\n".join(blocks)
    warnings: list[str] = []
    if not text.strip():
        raise ExtractionError("Word 文档里没有可提取的文字。图片和文本框暂不支持，请先转为文本。")

    # Word 的表格和段落已经保留换行；这里只清理过多空行。
    text = re.sub(r"\n{4,}", "\n\n\n", text).strip()
    return ExtractedDocument(
        text=text,
        source_type="docx",
        title=_clean_title(filename),
        warnings=warnings,
    )


_OCR_ENGINE = None
_OCR_LOCK = threading.Lock()


def _get_ocr():
    global _OCR_ENGINE
    with _OCR_LOCK:
        if _OCR_ENGINE is None:
            from rapidocr_onnxruntime import RapidOCR

            _OCR_ENGINE = RapidOCR()
        return _OCR_ENGINE


def _ocr_join_line(items: list[str]) -> str:
    result = ""
    for item in items:
        item = item.strip()
        if not item:
            continue
        if result and re.search(r"[A-Za-z0-9]$", result) and re.match(r"^[A-Za-z0-9]", item):
            result += " "
        result += item
    return result


def _ocr_to_text(result, warnings: list[str]) -> str:
    if not result:
        return ""
    rows = []
    for item in result:
        try:
            box, text, score = item
        except (TypeError, ValueError):
            continue
        if not text or float(score) < 0.38:
            continue
        points = np.array(box, dtype=float)
        x0, y0 = float(points[:, 0].min()), float(points[:, 1].min())
        x1, y1 = float(points[:, 0].max()), float(points[:, 1].max())
        rows.append(
            {
                "text": str(text).strip(),
                "score": float(score),
                "x0": x0,
                "x1": x1,
                "cy": (y0 + y1) / 2,
                "h": max(1.0, y1 - y0),
            }
        )
    if not rows:
        warnings.append("OCR 没有识别到可靠文字，请换一张更清晰的图片。")
        return ""

    rows.sort(key=lambda row: (row["cy"], row["x0"]))
    lines: list[list[dict]] = []
    for row in rows:
        target = None
        for line in reversed(lines[-3:]):
            reference = sum(item["cy"] for item in line) / len(line)
            tolerance = max(8.0, max(item["h"] for item in line) * 0.65)
            if abs(row["cy"] - reference) <= tolerance:
                target = line
                break
        if target is None:
            lines.append([row])
        else:
            target.append(row)

    output: list[str] = []
    last_cy = None
    last_h = None
    for line in lines:
        line.sort(key=lambda row: row["x0"])
        text = _ocr_join_line([row["text"] for row in line])
        if not text:
            continue
        cy = sum(row["cy"] for row in line) / len(line)
        h = max(row["h"] for row in line)
        if output and last_cy is not None and last_h is not None and cy - last_cy > max(24.0, h * 1.9):
            output.append("")
        output.append(text)
        last_cy = cy
        last_h = h
    return "\n".join(output)


def ocr_image(image: Image.Image, warnings: list[str] | None = None) -> str:
    warnings = warnings if warnings is not None else []
    image = ImageOps.exif_transpose(image).convert("RGB")
    max_side = max(image.size)
    if max_side > 2800:
        ratio = 2800 / max_side
        image = image.resize((max(1, int(image.width * ratio)), max(1, int(image.height * ratio))), Image.Resampling.LANCZOS)
    array = np.array(image)
    try:
        engine = _get_ocr()
        result, _ = engine(array)
    except Exception as exc:
        raise ExtractionError(f"OCR 引擎运行失败：{exc}") from exc
    return _ocr_to_text(result, warnings)


def _pdf_block_text(block: str) -> str:
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    if not lines:
        return ""
    output = lines[0]
    for line in lines[1:]:
        if re.search(r"[A-Za-z0-9]$", output) and re.match(r"^[A-Za-z0-9]", line):
            output += " "
        output += line
    return output


def extract_pdf(data: bytes, filename: str) -> ExtractedDocument:
    import pymupdf

    warnings: list[str] = []
    try:
        document = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise ExtractionError("PDF 文件无法解析，请确认文件没有损坏。") from exc

    if document.needs_pass:
        document.authenticate("")
    if document.needs_pass:
        document.close()
        raise ExtractionError("这个 PDF 有密码保护，请先解除密码后再上传。")

    if document.page_count > 80:
        document.close()
        raise ExtractionError("PDF 页数超过 80 页，请拆分成小文件后再上传。")

    page_texts: list[str] = []
    ocr_pages = 0
    for page_index in range(document.page_count):
        page = document[page_index]
        blocks = page.get_text("blocks", sort=True)
        text_blocks = []
        for block in blocks:
            if len(block) >= 5 and block[4] and block[4].strip():
                cleaned = _pdf_block_text(block[4])
                if cleaned:
                    text_blocks.append(cleaned)
        page_text = "\n\n".join(text_blocks).strip()
        if len(re.sub(r"\s", "", page_text)) < 12:
            # Scanned page or a page with only an image: render it and OCR locally.
            try:
                pix = page.get_pixmap(matrix=pymupdf.Matrix(2.0, 2.0), alpha=False)
                image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                ocr_text = ocr_image(image, warnings)
                if ocr_text.strip():
                    page_text = ocr_text.strip()
                    ocr_pages += 1
            except Exception as exc:
                warnings.append(f"第 {page_index + 1} 页 OCR 失败：{exc}")
        if page_text:
            page_texts.append(page_text)
        page = None

    document.close()
    text = "\n\n".join(page_texts).strip()
    if not text:
        raise ExtractionError("PDF 里没有提取到文字；如果是扫描件，请确保扫描清晰、方向正确。")
    if ocr_pages:
        warnings.append(f"检测到扫描版 PDF，已对 {ocr_pages} 页执行本地 OCR，请核对识别结果。")
    return ExtractedDocument(
        text=text,
        source_type="pdf",
        title=_clean_title(filename),
        warnings=warnings,
        pages=len(page_texts),
    )


def extract_image(data: bytes, filename: str) -> ExtractedDocument:
    try:
        image = Image.open(BytesIO(data))
        image.load()
    except (UnidentifiedImageError, OSError) as exc:
        raise ExtractionError("图片无法读取，请上传 PNG、JPG 或 WebP 图片。") from exc
    if image.width * image.height > 40_000_000:
        raise ExtractionError("图片像素过大，请压缩到 4000 万像素以内。")
    warnings: list[str] = []
    text = ocr_image(image, warnings)
    if not text.strip():
        raise ExtractionError("图片里没有识别到文字，请换一张更清晰、方向正确的图片。")
    warnings.append("图片文字由本地 OCR 识别，可能有少量错字，请在生成前核对。")
    return ExtractedDocument(
        text=text,
        source_type="image",
        title=_clean_title(filename),
        warnings=warnings,
        pages=1,
    )


def extract_text(data: bytes, filename: str) -> ExtractedDocument:
    text = _decode_text(data)
    if not text.strip():
        raise ExtractionError("文本文件是空的。")
    return ExtractedDocument(
        text=text,
        source_type="txt",
        title=_clean_title(filename),
    )


def extract_bytes(filename: str, data: bytes, content_type: str = "") -> ExtractedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt" or content_type.startswith("text/plain"):
        return extract_text(data, filename)
    if suffix == ".docx":
        return extract_docx(data, filename)
    if suffix == ".doc":
        raise ExtractionError("旧版 .doc 需要先另存为 .docx 再上传。")
    if suffix == ".pdf" or content_type == "application/pdf":
        if not data.startswith(b"%PDF"):
            raise ExtractionError("文件内容不是有效的 PDF。")
        return extract_pdf(data, filename)
    if suffix in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
        return extract_image(data, filename)
    raise ExtractionError("支持 .txt、.docx、.pdf、.png、.jpg、.jpeg、.webp 文件。")
